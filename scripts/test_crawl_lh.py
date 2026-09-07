#!/usr/bin/env python3
"""Unit tests for LH housing-list parsing helpers."""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import Workbook

from crawl_lh import (
    is_housing_list_file,
    normalize_housing_excel,
    notice_filename,
    parse_list_html,
    pick_housing_attachment,
    split_sido_sigungu,
)

LIST_HTML = """
<p class="bbs_total">전체 <strong>2 </strong>건 &nbsp;<strong>1</strong>/1페이지</p>
<div class="bbs_ListA">
<table>
<tbody>
<tr>
<td>2</td>
<td class="mVw cate col1">매입임대</td>
<td class="mVw bbs_tit">
<a href="javascript:" data-id1="2015122300020702" data-id2="03" data-id3="13" data-id4="26" class="wrtancInfoBtn">
<span>[경기남부] 용인시 신혼 매입임대 <em class="day">1일전</em></span>
</a>
</td>
<td class="mVw cate col2">경기도</td>
<td class="mVw"></td>
<td>2026.09.07</td>
<td>2026.09.21</td>
<td class="mVw stt noti">공고중</td>
<td>1</td>
</tr>
</tbody>
</table>
</div>
"""


class SplitAddressTests(unittest.TestCase):
    def test_seoul_district(self) -> None:
        sido, sigungu = split_sido_sigungu("서울특별시 강남구 논현로12길 23-5")
        self.assertEqual(sido, "서울특별시")
        self.assertEqual(sigungu, "강남구")

    def test_gyeonggi_city(self) -> None:
        sido, sigungu = split_sido_sigungu(
            "경기도 용인시 수지구 신수로783번길 9",
            fallback_sigungu="용인시",
        )
        self.assertEqual(sido, "경기도")
        self.assertEqual(sigungu, "용인시")


class ListParseTests(unittest.TestCase):
    def test_parse_list_html(self) -> None:
        total, notices = parse_list_html(LIST_HTML)
        self.assertEqual(total, 2)
        self.assertEqual(len(notices), 1)
        self.assertEqual(notices[0]["pan_id"], "2015122300020702")
        self.assertEqual(notices[0]["type"], "매입임대")
        self.assertIn("용인시 신혼 매입임대", notices[0]["title"])
        self.assertNotIn("1일전", notices[0]["title"])


class AttachmentTests(unittest.TestCase):
    def test_pick_posted_housing_xlsx(self) -> None:
        chosen = pick_housing_attachment(
            [
                {"cmnAhflNm": "공고문.pdf", "cmnAhflSn": 1},
                {"cmnAhflNm": "내부용 주택목록.xlsx", "cmnAhflSn": 2},
                {"cmnAhflNm": "★(게시용)주택목록.xlsx", "cmnAhflSn": 3},
            ]
        )
        self.assertIsNotNone(chosen)
        assert chosen is not None
        self.assertEqual(chosen["cmnAhflSn"], 3)

    def test_korean_housing_list_name(self) -> None:
        self.assertTrue(is_housing_list_file("★(게시용)용인시 주택목록(80호).xlsx"))
        self.assertTrue(is_housing_list_file("공급주택목록_게시용.xlsx"))
        self.assertTrue(is_housing_list_file("든든전세주택_주택보유목록.xlsx"))


class NormalizeExcelTests(unittest.TestCase):
    def test_multirow_header(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["공급대상 주택내역"])
        sheet.append(["순번", "지자체", "주택군명", "주택정보", "", "", "기본임대조건", ""])
        sheet.append(["", "", "", "주소", "동", "전용면적", "임대보증금", "월임대료"])
        sheet.append(
            [
                1,
                "용인시",
                "용인동천",
                "경기도 용인시 수지구 신수로 9",
                "101동",
                51.11,
                215600000,
                296240,
            ]
        )
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "sample.xlsx"
            workbook.save(path)
            frame = normalize_housing_excel(path)
        self.assertEqual(len(frame), 1)
        self.assertEqual(frame.iloc[0]["시도"], "경기도")
        self.assertEqual(frame.iloc[0]["시군구"], "용인시")
        self.assertEqual(frame.iloc[0]["보증금"], 215600000)
        self.assertEqual(frame.iloc[0]["전용면적"], 51.11)
        self.assertEqual(frame.iloc[0]["월임대료"], 296240)
        self.assertEqual(frame.iloc[0]["주택군"], "용인동천")

    def test_notice_filename(self) -> None:
        name = notice_filename("2026.09.07", "[경기남부] 용인시 신혼/매입임대")
        self.assertTrue(name.startswith("2026.09.07 LH "))
        self.assertNotIn("/", name)
        self.assertTrue(name.endswith(".xlsx"))
        spaced = notice_filename("2026.09.04", "[정정공고][경기남부] 청년 매입임대")
        self.assertIn("정정공고 경기남부", spaced)
        doubled = notice_filename(
            "2026.08.31",
            "[정정공고][정정공고][경북 포항,경주,영천] 일반 매입임대 입주자 모집공고",
        )
        self.assertIn("정정공고 경북", doubled)
        self.assertNotIn("정정공고 정정공고", doubled)
        glued = notice_filename("2026.08.27", "[인천지역본부]26년 6차 청년")
        self.assertIn("인천지역본부 26년", glued)

    def test_road_address_header(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["연번", "모집단위(주택군)", "소재지", "", "", "주택정보", ""])
        sheet.append(["", "", "광역지자체", "기초지자체", "도로명주소", "전용면적", "기본임대보증금"])
        sheet.append(
            [
                1,
                "26-6 남동구",
                "인천광역시",
                "남동구",
                "인천광역시 남동구 구월남로 116",
                82.44,
                15665000,
            ]
        )
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "incheon.xlsx"
            workbook.save(path)
            frame = normalize_housing_excel(path)
        self.assertEqual(len(frame), 1)
        self.assertEqual(frame.iloc[0]["시도"], "인천광역시")
        self.assertEqual(frame.iloc[0]["시군구"], "남동구")
        self.assertEqual(frame.iloc[0]["주소"], "인천광역시 남동구 구월남로 116")

    def test_deposit_on_sub_header_row(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["연번", "모집단위(주택군)", "소재지", "", "", "주택정보", "", "임대조건", ""])
        sheet.append(["", "", "광역지자체", "기초지자체", "도로명주소", "전용면적", "방수", "청년 1순위", ""])
        sheet.append(["", "", "", "", "", "", "", "기본임대보증금", "기본월임대료"])
        sheet.append(
            [
                1,
                "26-6 남동구",
                "인천광역시",
                "남동구",
                "인천광역시 남동구 구월남로 116",
                20.18,
                1,
                1_000_000,
                224690,
            ]
        )
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "incheon-sub.xlsx"
            workbook.save(path)
            frame = normalize_housing_excel(path)
        self.assertEqual(len(frame), 1)
        self.assertEqual(frame.iloc[0]["시도"], "인천광역시")
        self.assertEqual(frame.iloc[0]["보증금"], 1_000_000)
        self.assertEqual(frame.iloc[0]["월임대료"], 224690)
        self.assertEqual(frame.iloc[0]["주택구조(방수)"], 1)


if __name__ == "__main__":
    unittest.main()
